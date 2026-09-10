"""Excel tools — read / write / append xlsx workbooks (the classic RPA data source).

Three focused tools so a flow node says exactly what it does:
``excel.read``, ``excel.write``, ``excel.append``.

Requires the ``excel`` extra (``openpyxl``)::

    pip install "smithy-engine[excel]"

Paths honor the ``SMITHY_FILE_ROOT`` sandbox (see :mod:`smithy.core.files`).
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any

from smithy.core.blocking import run_blocking
from smithy.core.errors import InvalidInput, PlatformError
from smithy.core.files import confine_path
from smithy.core.tool import AbstractTool


class _ExcelBase(AbstractTool):
    """Shared ``path``/``sheet`` parsing for the Excel tools."""

    def _path(self, config: dict[str, Any]) -> Path:
        raw_path = config.get("path")
        if not isinstance(raw_path, (str, Path)) or not str(raw_path):
            raise InvalidInput(
                "Missing required parameter: path (expected an .xlsx path)",
                param="path",
                input_value=raw_path,
            )
        path = confine_path(Path(raw_path))
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        return path

    def _sheet(self, config: dict[str, Any]) -> str | None:
        sheet = config.get("sheet")
        if sheet is not None and (not isinstance(sheet, str) or not sheet):
            raise InvalidInput(
                "Invalid 'sheet': expected a non-empty string",
                param="sheet",
                input_value=sheet,
            )
        return sheet


class ExcelReadTool(_ExcelBase):
    """Read rows (and the header) from a workbook."""

    @property
    def name(self) -> str:
        return "excel.read"

    @property
    def description(self) -> str:
        return "Reads rows from an xlsx workbook"

    def schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Workbook path (.xlsx)"},
                "sheet": {
                    "type": "string",
                    "description": "Sheet name (default: active sheet)",
                },
                "header": {
                    "type": "boolean",
                    "default": True,
                    "description": "Treat the first row as a header",
                },
                "max_rows": {
                    "type": "integer",
                    "minimum": 1,
                    "default": 1000,
                    "description": "Max data rows to read",
                },
            },
            "required": ["path"],
        }

    async def execute(self, config: dict[str, Any]) -> Any:
        path = self._path(config)
        sheet = self._sheet(config)
        header = config.get("header", True)
        if not isinstance(header, bool):
            raise InvalidInput(
                "Invalid 'header': expected a boolean", param="header", input_value=header
            )
        max_rows = config.get("max_rows", 1000)
        if isinstance(max_rows, bool) or not isinstance(max_rows, int) or max_rows < 1:
            raise InvalidInput(
                "Invalid 'max_rows': expected an integer >= 1",
                param="max_rows",
                input_value=max_rows,
            )
        try:
            columns, rows = await run_blocking(_read_sync, path, sheet, header, max_rows)
        except FileNotFoundError as exc:
            raise PlatformError(f"Workbook not found: {path}", source=exc) from exc
        except PlatformError:
            raise
        except Exception as exc:
            raise PlatformError(f"Cannot read workbook {path}: {exc}", source=exc) from exc
        result: dict[str, Any] = {"rows": rows, "count": len(rows), "path": str(path)}
        if columns is not None:
            result["columns"] = columns
        return result


class ExcelWriteTool(_ExcelBase):
    """Write rows to a **new** workbook (refuses to replace by default)."""

    @property
    def name(self) -> str:
        return "excel.write"

    @property
    def description(self) -> str:
        return "Writes rows to a new xlsx workbook (refuses to overwrite by default)"

    def schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Workbook path (.xlsx)"},
                "sheet": {
                    "type": "string",
                    "description": "Sheet name (default: active sheet)",
                },
                "rows": {
                    "type": "array",
                    "items": {"type": "array"},
                    "description": "Rows to write (list of lists)",
                },
                "overwrite": {
                    "type": "boolean",
                    "default": False,
                    "description": "Allow write to replace an existing workbook",
                },
            },
            "required": ["path", "rows"],
        }

    async def execute(self, config: dict[str, Any]) -> Any:
        path = self._path(config)
        sheet = self._sheet(config)
        rows = _check_rows(config)
        overwrite = config.get("overwrite", False)
        if not isinstance(overwrite, bool):
            raise InvalidInput(
                "Invalid 'overwrite': expected a boolean",
                param="overwrite",
                input_value=overwrite,
            )
        written = await run_blocking(_write_sync, path, sheet, rows, overwrite)
        return {"path": str(path), "written": written}


class ExcelAppendTool(_ExcelBase):
    """Append rows to an existing workbook."""

    @property
    def name(self) -> str:
        return "excel.append"

    @property
    def description(self) -> str:
        return "Appends rows to an existing xlsx workbook"

    def schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Workbook path (.xlsx)"},
                "sheet": {
                    "type": "string",
                    "description": "Sheet name (default: active sheet)",
                },
                "rows": {
                    "type": "array",
                    "items": {"type": "array"},
                    "description": "Rows to append (list of lists)",
                },
            },
            "required": ["path", "rows"],
        }

    async def execute(self, config: dict[str, Any]) -> Any:
        path = self._path(config)
        sheet = self._sheet(config)
        rows = _check_rows(config)
        try:
            written = await run_blocking(_append_sync, path, sheet, rows)
        except FileNotFoundError as exc:
            raise PlatformError(f"Workbook not found: {path}", source=exc) from exc
        except PlatformError:
            raise
        except Exception as exc:
            raise PlatformError(f"Cannot append to workbook {path}: {exc}", source=exc) from exc
        return {"path": str(path), "written": written}


def _write_sync(path: Path, sheet: str | None, rows: list[list[Any]], overwrite: bool) -> int:
    """Write a fresh workbook (runs in an executor)."""
    if path.exists() and not overwrite:
        raise PlatformError(f"Workbook already exists: {path} (pass overwrite=true to replace)")
    from openpyxl import Workbook

    try:
        workbook = Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        if sheet is not None:
            # Rename the default sheet instead of adding a second one, so the
            # workbook does not carry an empty "Sheet".
            worksheet.title = sheet
        for row in rows:
            worksheet.append(row)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(path))
    except PlatformError:
        raise
    except Exception as exc:
        raise PlatformError(f"Cannot write workbook {path}: {exc}", source=exc) from exc
    return len(rows)


def _read_sync(
    path: Path, sheet: str | None, header: bool, max_rows: int
) -> tuple[list[str] | None, list[list[Any]]]:
    """Read rows (runs in an executor)."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True, read_only=True)
    try:
        worksheet = workbook[sheet] if sheet is not None else workbook.active
        if worksheet is None:
            raise PlatformError(f"Workbook {path} has no sheets")
        columns: list[str] | None = None
        rows: list[list[Any]] = []
        for index, row in enumerate(worksheet.iter_rows(values_only=True)):
            values = [_cell(value) for value in row]
            if index == 0 and header:
                columns = [str(value) for value in values]
                continue
            rows.append(values)
            if len(rows) >= max_rows:
                break
        return columns, rows
    finally:
        workbook.close()


def _append_sync(path: Path, sheet: str | None, rows: list[list[Any]]) -> int:
    """Append rows to an existing workbook (runs in an executor)."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path))
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            raise PlatformError(f"Sheet {sheet!r} not found in {path}")
        worksheet = workbook[sheet] if sheet is not None else workbook.active
        assert worksheet is not None
        for row in rows:
            worksheet.append(row)
        workbook.save(str(path))
        return len(rows)
    finally:
        workbook.close()


def _check_rows(config: dict[str, Any]) -> list[list[Any]]:
    raw_rows = config.get("rows")
    if not isinstance(raw_rows, list) or not all(isinstance(row, list) for row in raw_rows):
        raise InvalidInput(
            "Missing or invalid 'rows': expected a list of lists",
            param="rows",
            input_value=type(raw_rows).__name__ if not isinstance(raw_rows, list) else raw_rows,
        )
    return raw_rows


def _cell(value: Any) -> Any:
    """JSON-safe cell value (runs in an executor context, never raises)."""
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.isoformat()
    return value
